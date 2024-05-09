from openpyxl import load_workbook
import json

RANGE = [50, 50]

NAME_KEY = 'N'
EMAIL_KEY = 'E'
CHAMPION_KEY = 'FP'
LOCAL_TEAM_KEY = 'LT'
VISIT_TEAM_KEY = 'VT'
LOCAL_SCORE_KEY = 'LTS'
VISIT_SCORE_KEY = 'VTS'

def _search_for_value(workbook, value):
    
    # Iterate through all sheets in the workbook
    for sheet_name in workbook.sheetnames:
        # Get the current sheet
        sheet = workbook[sheet_name]
        
        # Iterate through all rows and columns in the sheet
        iter_rows = 0
        for row in sheet.iter_rows():
            iter_cell = 0
            for cell in row:
                # Check if the cell value matches the value we're searching for
                if cell.value == value:
                    # Return the coordinate of the cell where the value is found
                    return [sheet_name, str(cell.coordinate)]
                iter_cell += 1
            iter_rows += 1
    
    # If the value is not found, return None
    return None

def generate_ranges(filename, games = [], save_folder = ''):
    # Open workbook
    workbook = load_workbook(filename=filename, data_only = True)
    games_ranges = {}
    for g in games:
        games_ranges[g] = {
            'data': {
                'local_team': _search_for_value(workbook, "{}{}".format(LOCAL_TEAM_KEY, g)),
                'visit_team': _search_for_value(workbook, "{}{}".format(VISIT_TEAM_KEY, g)),
                'local_score': _search_for_value(workbook, "{}{}".format(LOCAL_SCORE_KEY, g)),
                'visit_score': _search_for_value(workbook, "{}{}".format(VISIT_SCORE_KEY, g))
            }
        }
    ranges = {
        "name": _search_for_value(workbook, NAME_KEY),
        "email": _search_for_value(workbook, EMAIL_KEY),
        "champion": _search_for_value(workbook, CHAMPION_KEY),
        "games": games_ranges
    }
    with open('{}/ranges.json'.format(save_folder), "w") as file:
        json.dump(ranges, file)