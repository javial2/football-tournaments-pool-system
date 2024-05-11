from openpyxl import load_workbook
import json

def generate_configuration(filename, save_folder = ''):
    # Open workbook
    workbook = load_workbook(filename=filename, data_only = True)
    configuration_sheet = workbook['Tournament Configuration']
    configuration = {row[0].value: row[1].value for row in configuration_sheet['B2:C4']}
    tables = configuration_sheet.tables
    stages = [s[0].value for s in configuration_sheet[tables['Table7'].ref][1:]]
    games = {str(r[0].value): r[1].value for r in configuration_sheet[tables['Table8'].ref][1:]}
    
    tournament_configuration = {
        "tournament_name": configuration['Tournament Name'],
        "players_database": "database/players_database.json",
        "results_database": "database/results_database.json",
        "stats_folder": "stats/",
        "results_xlsx": configuration["Results Filename"],
        "ranges_xlsx": configuration["Ranges Filename"],
        "tournament_stages": {
            format_names(s): {
                "nid": format_names(s),
                "order": stages.index(s) + 1,
                "games": [g for g in games.keys() if s == games[g]]
            } for s in stages
        }
    }

    with open('{}/tournament_configuration.json'.format(save_folder), "w") as file:
        json.dump(tournament_configuration, file)
    return tournament_configuration

def format_names(text=''):
    text = text.lower()
    text = text.replace(' ', '-')
    text = text.replace('á', 'a')
    text = text.replace('é', 'e')
    text = text.replace('í', 'i')
    text = text.replace('ó', 'o')
    text = text.replace('ú', 'u')
    
    return text 

#generate_configuration('instances/Test Mundial Qatar 2022/CONFIGURATION.xlsx', save_folder='instances/Test Mundial Qatar 2022/config')