from openpyxl import load_workbook
import json

#from src.utils.utils import format_names

DICTIONARY = {
    'Champion': 'champion',
    'Top Goalscorer': 'top_goalscorer',
    'Qualified Teams': 'teams',
    'Result': 'points_per_result',
    'Score': 'points_per_score',
    'Teams': 'points_per_team',
    'FALSE': False,
    'TRUE': True
}

PUNCTUATION_SYSTEMS = {
    'tournament': ['champion', 'top_goalscorer'],
    'stage': ['teams'],
    'game': ['points_per_result', 'points_per_score', 'points_per_team']
}

RESTRICTION_SYSTEMS = {
    'games': ['teams', 'reversible', 'result']
}

def generate_points_configuration(filename, stages = [], save_folder = ''):
    # Open workbook
    workbook = load_workbook(filename=filename, data_only = True)
    points_sheet = workbook['Points System']
    tables = points_sheet.tables
    points = [[c.value for c in r] for r in points_sheet[tables['Table6'].ref][1:]]

    points_configuration = {
        'tournament': {
            'stages': {
                s: {
                    'games': {}
                } for s in stages
            }
        }
    }

    for row in points:
        point_system = DICTIONARY[row[0]]
        for ps in PUNCTUATION_SYSTEMS:
            if point_system in PUNCTUATION_SYSTEMS[ps]:
                dependance = ps
                break
        if dependance == 'tournament':
            points_configuration['tournament'][point_system] = {
                'value': row[2],
                'restrictions': {}
            }
        elif dependance == 'stage':
            stage = format_names(row[1])
            points_configuration['tournament']['stages'][stage][point_system] = {
                'value': row[2],
                'restrictions': {}
            }
        elif dependance == 'game':
            stage = format_names(row[1])
            points_configuration['tournament']['stages'][stage]['games'][point_system] = {
                'value': row[2],
                'restrictions': {
                    'teams': row[3],
                    'reversible': row[4],
                    'result': row[5]
                }
            }
    
    with open('{}/points.json'.format(save_folder), "w") as file:
        json.dump(points_configuration, file)
    return points_configuration

def format_names(text=''):
    text = text.lower()
    text = text.replace(' ', '-')
    text = text.replace('á', 'a')
    text = text.replace('é', 'e')
    text = text.replace('í', 'i')
    text = text.replace('ó', 'o')
    text = text.replace('ú', 'u')
    
    return text        

#generate_points_configuration('instances/Test Mundial Qatar 2022/CONFIGURATION.xlsx', stages = ['groups', 'top-16', 'top-8', 'top-4', 'finals', 'third-place'], save_folder='instances/Test Mundial Qatar 2022/config')