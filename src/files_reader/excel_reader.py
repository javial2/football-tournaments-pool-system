from openpyxl import load_workbook
import json
from copy import deepcopy

def detect_excel_range(workbook, element):
    if isinstance(element, list):
        if len(element) == 2:
            if element[0] in workbook.sheetnames:
                return True
    return False

def element_iterator(workbook, element):
    if detect_excel_range(workbook, element):
        return workbook[element[0]][element[1]].value
    elif isinstance(element, list):
        key = 0
        while key < len(element):
            value = element[key]
            element[key] = element_iterator(workbook, value)
            key += 1
    elif isinstance(element, dict):
        for key in element:
            value = element[key]
            element[key] = element_iterator(workbook, value)
    return element

def excel_file_reader(filename, ranges):
    workbook = load_workbook(filename=filename, data_only = True)
    #with open(ranges) as json_file:
        #ranges_file = json.load(json_file)
    ranges_file = deepcopy(ranges)
    for k in ranges_file:
        e = ranges_file[k]
        ranges_file[k] = element_iterator(workbook, e)
    return ranges_file
    

#excel_file_reader("././instances/Copa America 2021 Familiar/RESULTADOS_REALES.xlsx", "././instances/Copa America 2021 Familiar/config/ranges.json")
